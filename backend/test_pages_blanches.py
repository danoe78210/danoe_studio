import os
import sys
import platform

import pytest

sys.path.insert(0, os.path.dirname(__file__))

try:
    import win32com.client  # noqa
    PYWIN32_DISPONIBLE = True
except ImportError:
    PYWIN32_DISPONIBLE = False

from docx.enum.section import WD_SECTION_START

import generer_roman
from generer_roman import corriger_doubles_pages_blanches, compter_pages_reelles, fermer_word


@pytest.mark.skipif(
    platform.system() != 'Windows' or not PYWIN32_DISPONIBLE,
    reason='Nécessite Word + pywin32, disponible uniquement sous Windows',
)
def test_corriger_doubles_pages_blanches_fusionne_une_page(tmp_path):
    # Réutilise le document global du module (créé à l'import, sans contenu)
    # ainsi que sa fonction nouvelle_section, pour reproduire fidèlement le
    # pattern réel de generer_roman.inserer_blanks() : une section
    # « nouvelle page » + un paragraphe vide, pour CHAQUE page blanche.
    document = generer_roman.doc
    document.add_paragraph('Page avec du texte, avant les pages blanches.')

    for _ in range(2):
        sb = generer_roman.nouvelle_section(WD_SECTION_START.NEW_PAGE, 'top')
        generer_roman.definir_entete(sb, '')
        generer_roman.definir_pieds(sb, False)
        document.add_paragraph()

    # Comme dans saut_vers_page_impaire() : une nouvelle section est toujours
    # ouverte après les pages blanches, avant le contenu suivant, pour que
    # celui-ci démarre sur sa propre nouvelle page.
    sec_apres = generer_roman.nouvelle_section(WD_SECTION_START.NEW_PAGE, 'top')
    generer_roman.definir_entete(sec_apres, '')
    generer_roman.definir_pieds(sec_apres, False)
    document.add_paragraph('Page avec du texte, après les pages blanches.')

    chemin_docx = str(tmp_path / 'test_pages_blanches.docx')
    document.save(chemin_docx)

    try:
        pages_avant = compter_pages_reelles(chemin_docx)
        assert pages_avant is not None
        # Garantit que le cas testé est bien réel : texte + blanche + blanche + texte
        assert pages_avant == 4

        corriger_doubles_pages_blanches(chemin_docx)

        pages_apres = compter_pages_reelles(chemin_docx)
        assert pages_apres is not None
        assert pages_apres == pages_avant - 1
    finally:
        try:
            fermer_word()
        except Exception:
            pass
